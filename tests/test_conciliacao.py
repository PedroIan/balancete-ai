"""
Testes da conciliação bancária em core/conciliacao.py
Execute com: python tests/test_conciliacao.py
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

import io

import openpyxl

from core.conciliacao import conciliar, gerar_xlsx


def test_conciliacao_basica():
    extrato = [
        {"data": "2026-05-10", "descricao": "PAG CEMIG", "valor": 450.00, "tipo": "debito", "saldo": None, "fonte": "extrato"},
        {"data": "2026-05-12", "descricao": "PAG COPASA", "valor": 120.00, "tipo": "debito", "saldo": None, "fonte": "extrato"},
        {"data": "2026-05-15", "descricao": "TRANSF PIX", "valor": 800.00, "tipo": "debito", "saldo": None, "fonte": "extrato"},
    ]
    transacoes = [
        {"data": "2026-05-10", "descricao": "Energia elétrica", "fornecedor": "CEMIG", "valor": 450.00, "tipo": "despesa", "categoria": "Energia", "cnpj": "", "fonte": "nf1.pdf"},
        {"data": "2026-05-11", "descricao": "Água e esgoto", "fornecedor": "COPASA", "valor": 120.00, "tipo": "despesa", "categoria": "Água", "cnpj": "", "fonte": "nf2.pdf"},
    ]

    pares, ext_sem, tx_sem = conciliar(extrato, transacoes)

    assert len(pares) == 2, f"Esperava 2 pares, obteve {len(pares)}"
    assert len(ext_sem) == 1
    assert ext_sem[0]["valor"] == 800.00
    assert len(tx_sem) == 0

    print("✅ test_conciliacao_basica passou")


def test_conciliacao_tolerancia():
    extrato = [
        {"data": "2026-05-10", "descricao": "BOLETO", "valor": 1000.03, "tipo": "debito", "saldo": None, "fonte": "extrato"},
    ]
    transacoes = [
        {"data": "2026-05-12", "descricao": "Pagamento", "fornecedor": "Fornecedor X", "valor": 1000.00, "tipo": "despesa", "categoria": "Outras", "cnpj": "", "fonte": "comprov.pdf"},
    ]
    pares, ext_sem, tx_sem = conciliar(extrato, transacoes)
    # Valor dentro da tolerância (diff 0.03) e data dentro de 3 dias (diff 2)
    assert len(pares) == 1, f"Esperava 1 par por tolerância, obteve {len(pares)}"

    print("✅ test_conciliacao_tolerancia passou")


def test_conciliacao_sem_par():
    extrato = [
        {"data": "2026-05-10", "descricao": "TRANSF", "valor": 999.00, "tipo": "debito", "saldo": None, "fonte": "extrato"},
    ]
    transacoes = [
        {"data": "2026-05-10", "descricao": "Pagamento", "fornecedor": "X", "valor": 500.00, "tipo": "despesa", "categoria": "Outras", "cnpj": "", "fonte": "doc.pdf"},
    ]
    pares, ext_sem, tx_sem = conciliar(extrato, transacoes)
    assert len(pares) == 0
    assert len(ext_sem) == 1
    assert len(tx_sem) == 1

    print("✅ test_conciliacao_sem_par passou")


def test_receita_suspeita_listada():
    """Receita suspeita deve aparecer na aba Receitas (entra no total) E na aba Revisar."""
    transacoes = [
        {"data": "2026-05-05", "descricao": "Taxa condominial 101", "fornecedor": "APTO 101",
         "valor": 600.00, "tipo": "receita", "categoria": "Condomínio", "cnpj": "",
         "suspeito": False, "fonte": "rec1.pdf"},
        {"data": "2026-05-06", "descricao": "", "fornecedor": "APTO 102",
         "valor": 600.00, "tipo": "receita", "categoria": "Condomínio", "cnpj": "",
         "suspeito": True, "fonte": "rec2.pdf"},
    ]
    xlsx = gerar_xlsx(transacoes, [], "2026-05", 1000.0)
    wb = openpyxl.load_workbook(io.BytesIO(xlsx))

    ws_rec = wb["Receitas"]
    # Cabeçalho + 2 receitas (inclusive a suspeita)
    assert ws_rec.max_row == 3, f"Esperava 3 linhas na aba Receitas, obteve {ws_rec.max_row}"

    assert "⚠️ Revisar" in wb.sheetnames

    # Total de receitas no Resumo inclui a suspeita: 600 + 600 = 1200
    ws_resumo = wb["Resumo"]
    valores = {ws_resumo.cell(r, 1).value: ws_resumo.cell(r, 2).value for r in range(2, 6)}
    assert valores["(+) Total Receitas"] == 1200.00
    assert valores["= Saldo Final"] == 2200.00

    print("✅ test_receita_suspeita_listada passou")


if __name__ == "__main__":
    test_conciliacao_basica()
    test_conciliacao_tolerancia()
    test_conciliacao_sem_par()
    test_receita_suspeita_listada()
    print("\n✅ Todos os testes de conciliação passaram")
